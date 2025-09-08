"""Report generation module for Excel and other outputs"""

import os
import json
import pandas as pd
import numpy as np
from typing import Dict, List, Optional, Any
from datetime import datetime, date
import joblib


class CustomJSONEncoder(json.JSONEncoder):
    """Custom JSON encoder for handling pandas/numpy types"""
    def default(self, obj):
        if isinstance(obj, (pd.Timestamp, datetime, date)):
            return obj.isoformat()
        elif isinstance(obj, pd.Series):
            return obj.tolist()
        elif isinstance(obj, np.ndarray):
            return obj.tolist()
        elif isinstance(obj, (np.integer, np.int64, np.int32)):
            return int(obj)
        elif isinstance(obj, (np.floating, np.float64, np.float32)):
            return float(obj)
        elif isinstance(obj, np.bool_):
            return bool(obj)
        elif pd.isna(obj):
            return None
        return super().default(obj)

class ReportGenerator:
    """Handles report generation and export"""
    
    def __init__(self, config):
        self.cfg = config
        self.sheets = {}
        
    def build_model_summary_report(
        self, 
        models_summary: pd.DataFrame,
        best_model_name: str
    ) -> Dict[str, pd.DataFrame]:
        """Build model summary reports"""
        reports = {}
        
        # Overall model summary
        if models_summary is not None:
            reports["models_summary"] = models_summary
            
            # Best model details
            if best_model_name:
                best_model_df = models_summary[
                    models_summary["model_name"] == best_model_name
                ]
                reports["best_model"] = best_model_df
        
        return reports
    
    def build_feature_importance_report(
        self, 
        model,
        feature_names: List[str],
        woe_map: Optional[Dict] = None
    ) -> pd.DataFrame:
        """Build feature importance report"""
        rows = []
        
        if hasattr(model, "coef_"):
            # Linear models
            coefs = model.coef_[0] if model.coef_.ndim > 1 else model.coef_
            
            # Ensure coefs and feature_names have same length
            min_len = min(len(coefs), len(feature_names))
            for i in range(min_len):
                rows.append({
                    "variable": feature_names[i],
                    "coef_or_importance": coefs[i],
                    "sign": np.sign(coefs[i]),
                    "variable_group": self._get_variable_group(feature_names[i], woe_map)
                })
                
        elif hasattr(model, "feature_importances_"):
            # Tree-based models
            importances = model.feature_importances_
            
            # Ensure importances and feature_names have same length
            min_len = min(len(importances), len(feature_names))
            for i in range(min_len):
                rows.append({
                    "variable": feature_names[i],
                    "coef_or_importance": importances[i],
                    "sign": 1 if importances[i] >= 0 else -1,
                    "variable_group": self._get_variable_group(feature_names[i], woe_map)
                })
        else:
            # Other models
            for feature in feature_names:
                rows.append({
                    "variable": feature,
                    "coef_or_importance": 0,
                    "sign": 0,
                    "variable_group": self._get_variable_group(feature, woe_map)
                })
        
        df = pd.DataFrame(rows)
        
        # Sort by absolute importance
        if not df.empty:
            df = df.sort_values(
                "coef_or_importance",
                key=lambda x: np.abs(x),
                ascending=False
            ).reset_index(drop=True)
        
        return df
    
    def _get_variable_group(self, variable: str, woe_map: Optional[Dict]) -> str:
        """Get variable group from WOE map"""
        if woe_map and variable in woe_map:
            return woe_map[variable].var_type
        return "unknown"
    
    def build_iv_report(self, woe_map: Dict, top_n: int = 20) -> pd.DataFrame:
        """Build Information Value report"""
        rows = []
        
        for var, vw in woe_map.items():
            rows.append({
                "variable": var,
                "IV": vw.iv,
                "variable_group": vw.var_type
            })
        
        df = pd.DataFrame(rows)
        
        if not df.empty:
            df = df.sort_values("IV", ascending=False).head(top_n).reset_index(drop=True)
        
        return df
    
    def build_psi_report(self, psi_results: pd.DataFrame) -> Dict[str, pd.DataFrame]:
        """Build PSI reports"""
        reports = {}
        
        if psi_results is not None and not psi_results.empty:
            reports["psi_summary"] = psi_results
            
            # Dropped features
            dropped = psi_results[psi_results["status"] == "DROP"]
            if not dropped.empty:
                reports["psi_dropped_features"] = dropped
        
        return reports
    
    def build_woe_mapping_report(self, woe_map: Dict) -> pd.DataFrame:
        """Build WOE mapping report"""
        rows = []
        
        for var, vw in woe_map.items():
            if vw.var_type == "numeric" and vw.numeric_bins:
                for bin_obj in vw.numeric_bins:
                    rows.append({
                        "variable": var,
                        "type": "numeric",
                        "item_type": "bin",
                        "left": bin_obj.left,
                        "right": bin_obj.right,
                        "label": None,
                        "members": None,
                        "woe": bin_obj.woe,
                        "event_rate": bin_obj.event_rate,
                        "iv_contrib": bin_obj.iv_contrib
                    })
            elif vw.var_type == "categorical" and vw.categorical_groups:
                for group in vw.categorical_groups:
                    rows.append({
                        "variable": var,
                        "type": "categorical",
                        "item_type": "group",
                        "left": None,
                        "right": None,
                        "label": group.label,
                        "members": ", ".join(map(str, group.members)) if group.members else None,
                        "woe": group.woe,
                        "event_rate": group.event_rate,
                        "iv_contrib": group.iv_contrib
                    })
        
        return pd.DataFrame(rows)
    
    def build_univariate_analysis(
        self, 
        X: pd.DataFrame, 
        y: np.ndarray,
        top_n: int = 50
    ) -> pd.DataFrame:
        """Build univariate analysis report"""
        from sklearn.feature_selection import f_classif
        
        rows = []
        
        for col in X.columns:
            x_col = X[col]
            
            # Skip if all missing
            if x_col.notna().sum() == 0:
                continue
            
            # Calculate basic stats
            corr = np.corrcoef(x_col.fillna(0), y)[0, 1]
            
            # F-statistic
            try:
                f_stat, p_val = f_classif(
                    x_col.fillna(0).values.reshape(-1, 1), 
                    y
                )
                f_stat = float(f_stat[0])
                p_val = float(p_val[0])
            except Exception:
                f_stat = 0.0
                p_val = 1.0
            
            rows.append({
                "variable": col,
                "correlation": corr,
                "f_statistic": f_stat,
                "p_value": p_val,
                "missing_pct": x_col.isna().mean()
            })
        
        df = pd.DataFrame(rows)
        
        if not df.empty:
            df = df.sort_values("f_statistic", ascending=False).head(top_n).reset_index(drop=True)
        
        return df
    
    def export_to_excel(self, output_path: str, sheets: Dict[str, pd.DataFrame]):
        """Export reports to Excel"""
        try:
            # Try with openpyxl first
            with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
                for sheet_name, df in sheets.items():
                    if df is not None and not df.empty:
                        # Truncate sheet name to 31 characters (Excel limit)
                        sheet_name = sheet_name[:31]
                        df.to_excel(writer, sheet_name=sheet_name, index=False)
                        
                        # Format as table
                        try:
                            self._format_as_table(writer, df, sheet_name)
                        except Exception:
                            pass
        except Exception:
            # Fallback to xlsxwriter
            try:
                with pd.ExcelWriter(output_path, engine="xlsxwriter") as writer:
                    for sheet_name, df in sheets.items():
                        if df is not None and not df.empty:
                            sheet_name = sheet_name[:31]
                            df.to_excel(writer, sheet_name=sheet_name, index=False)
            except Exception as e:
                print(f"Failed to export Excel: {e}")
    
    def _format_as_table(self, writer, df: pd.DataFrame, sheet_name: str):
        """Format Excel sheet as table"""
        try:
            worksheet = writer.sheets[sheet_name]
            nrows, ncols = df.shape
            
            # Try openpyxl formatting
            from openpyxl.utils import get_column_letter
            from openpyxl.worksheet.table import Table, TableStyleInfo
            
            ref = f"A1:{get_column_letter(ncols)}{nrows+1}"
            table = Table(displayName=f"tbl_{sheet_name}", ref=ref)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium9",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False
            )
            worksheet.add_table(table)
        except Exception:
            # Try xlsxwriter formatting
            try:
                worksheet.add_table(
                    0, 0, nrows, ncols-1,
                    {
                        "name": f"tbl_{sheet_name}",
                        "style": "Table Style Medium 9"
                    }
                )
            except Exception:
                pass
    
    def export_artifacts(
        self,
        output_folder: str,
        run_id: str,
        woe_map: Dict,
        final_vars: List[str],
        best_model,
        shap_values: Optional[Dict] = None
    ):
        """Export pipeline artifacts"""
        os.makedirs(output_folder, exist_ok=True)
        
        # Export WOE mapping
        woe_dict = {}
        for var, vw in woe_map.items():
            woe_dict[var] = vw.to_dict() if hasattr(vw, 'to_dict') else str(vw)
        
        with open(os.path.join(output_folder, f"woe_mapping_{run_id}.json"), "w") as f:
            json.dump(woe_dict, f, ensure_ascii=False, indent=2, cls=CustomJSONEncoder)
        
        # Export final variables
        with open(os.path.join(output_folder, f"final_vars_{run_id}.json"), "w") as f:
            json.dump({"final_vars": final_vars}, f, ensure_ascii=False, indent=2, cls=CustomJSONEncoder)
        
        # Export best model
        if best_model:
            joblib.dump(
                best_model,
                os.path.join(output_folder, f"best_model_{run_id}.joblib")
            )
        
        # Export SHAP values if available
        if shap_values:
            with open(os.path.join(output_folder, f"shap_summary_{run_id}.json"), "w") as f:
                json.dump(shap_values, f, ensure_ascii=False, indent=2)
    
    def generate_full_report(
        self,
        models_summary: pd.DataFrame,
        best_model_name: str,
        best_model,
        final_vars: List[str],
        woe_map: Dict,
        psi_results: Optional[pd.DataFrame] = None,
        X_train: Optional[pd.DataFrame] = None,
        y_train: Optional[np.ndarray] = None,
        shap_values: Optional[Any] = None,
        shap_summary: Optional[Dict] = None
    ) -> Dict[str, pd.DataFrame]:
        """Generate full report with all sheets"""
        sheets = {}
        
        # Model summary sheets
        model_reports = self.build_model_summary_report(models_summary, best_model_name)
        sheets.update(model_reports)
        
        # Feature importance
        if best_model and final_vars:
            sheets["best_model_vars"] = self.build_feature_importance_report(
                best_model, final_vars, woe_map
            )
        
        # IV report
        if woe_map:
            sheets["top20_iv"] = self.build_iv_report(woe_map, top_n=20)
        
        # PSI reports
        if psi_results is not None:
            psi_reports = self.build_psi_report(psi_results)
            sheets.update(psi_reports)
        
        # WOE mapping
        if woe_map:
            sheets["woe_mapping"] = self.build_woe_mapping_report(woe_map)
        
        # Comprehensive best model variables sheet
        if best_model and final_vars and woe_map:
            sheets["best_model_details"] = self.build_best_model_comprehensive_report(
                best_model, final_vars, woe_map, best_model_name
            )
        
        # SHAP analysis sheet
        if shap_values is not None and shap_summary is not None:
            sheets["shap_importance"] = self.build_shap_report(
                shap_values, shap_summary, final_vars
            )
        
        # Univariate analysis
        if X_train is not None and y_train is not None:
            sheets["top50_univariate"] = self.build_univariate_analysis(
                X_train, y_train, top_n=50
            )
        
        # Metadata
        sheets["run_meta"] = self._build_metadata_report()
        
        return sheets
    
    def _extract_feature_importance(self, model) -> Optional[np.ndarray]:
        """Extract feature importance from model"""
        if model is None:
            return None
            
        # Tree-based models
        if hasattr(model, 'feature_importances_'):
            return model.feature_importances_
        
        # Linear models
        elif hasattr(model, 'coef_'):
            coef = model.coef_
            if len(coef.shape) > 1:
                coef = coef[0]
            return np.abs(coef)
        
        # GAM models
        elif hasattr(model, 'statistics_'):
            if hasattr(model.statistics_, 'p_values'):
                # Use inverse of p-values as importance
                p_values = model.statistics_.p_values
                return 1.0 / (p_values + 0.001) if p_values is not None else None
        
        return None
    
    def build_best_model_comprehensive_report(
        self, 
        best_model, 
        final_vars: List[str], 
        woe_map: Dict,
        best_model_name: str
    ) -> pd.DataFrame:
        """Build comprehensive best model variables report with WOE details"""
        rows = []
        
        # Get feature importances
        importances = self._extract_feature_importance(best_model)
        importance_dict = dict(zip(final_vars, importances)) if importances is not None else {}
        
        # Check if this is a WOE or RAW model
        is_woe_model = best_model_name.startswith("WOE_")
        
        for var in final_vars:
            # Get base variable name (remove _woe suffix if present)
            base_var = var.replace("_woe", "") if var.endswith("_woe") else var
            
            # Get WOE mapping info if available
            if base_var in woe_map:
                vw = woe_map[base_var]
                
                # For numeric variables with bins
                if vw.var_type == "numeric" and vw.numeric_bins:
                    for i, bin_obj in enumerate(vw.numeric_bins):
                        rows.append({
                            "rank": 0,  # Will be updated after sorting
                            "variable": base_var,
                            "type": "numeric",
                            "bin_number": i + 1,
                            "bin_range": f"[{bin_obj.left:.2f}, {bin_obj.right:.2f})",
                            "woe": bin_obj.woe,  # Always show WOE values
                            "event_count": bin_obj.event_count,
                            "nonevent_count": bin_obj.nonevent_count,
                            "total_count": bin_obj.total_count,
                            "event_rate": bin_obj.event_rate,
                            "iv_contrib": bin_obj.iv_contrib,
                            "total_iv": vw.iv,
                            "importance": importance_dict.get(var, 0),
                            "model_type": "WOE" if is_woe_model else "RAW"
                        })
                
                # For categorical variables with groups
                elif vw.var_type == "categorical" and vw.categorical_groups:
                    for i, group in enumerate(vw.categorical_groups):
                        rows.append({
                            "rank": 0,
                            "variable": base_var,
                            "type": "categorical", 
                            "bin_number": i + 1,
                            "bin_range": f"{group.label}: {','.join(str(m) for m in group.members[:3])}...",
                            "woe": group.woe,  # Always show WOE values
                            "event_count": group.event_count,
                            "nonevent_count": group.nonevent_count,
                            "total_count": group.total_count,
                            "event_rate": group.event_rate,
                            "iv_contrib": group.iv_contrib,
                            "total_iv": vw.iv,
                            "importance": importance_dict.get(var, 0),
                            "model_type": "WOE" if is_woe_model else "RAW"
                        })
                
                # Add missing value row if exists
                if vw.missing_woe is not None and vw.missing_rate > 0:
                    rows.append({
                        "rank": 0,
                        "variable": base_var,
                        "type": vw.var_type,
                        "bin_number": 999,
                        "bin_range": "Missing",
                        "woe": vw.missing_woe if is_woe_model else None,
                        "event_count": None,
                        "nonevent_count": None,
                        "total_count": None,
                        "event_rate": vw.missing_rate,
                        "iv_contrib": None,
                        "total_iv": vw.iv,
                        "importance": importance_dict.get(var, 0),
                        "model_type": "WOE" if is_woe_model else "RAW"
                    })
            else:
                # Variable not in WOE map (RAW variable)
                rows.append({
                    "rank": 0,
                    "variable": var,
                    "type": "raw",
                    "bin_number": None,
                    "bin_range": "RAW feature",
                    "woe": None,
                    "event_count": None,
                    "nonevent_count": None,
                    "total_count": None,
                    "event_rate": None,
                    "iv_contrib": None,
                    "total_iv": None,
                    "importance": importance_dict.get(var, 0),
                    "model_type": "RAW"
                })
        
        df = pd.DataFrame(rows)
        
        if not df.empty:
            # Sort by importance (descending) and then by variable name
            df = df.sort_values(
                ["importance", "variable", "bin_number"],
                ascending=[False, True, True]
            ).reset_index(drop=True)
            
            # Update rank based on unique variables
            seen_vars = set()
            rank = 0
            for idx, row in df.iterrows():
                if row["variable"] not in seen_vars:
                    rank += 1
                    seen_vars.add(row["variable"])
                df.at[idx, "rank"] = rank
            
            # Reorder columns for better readability
            column_order = [
                "rank", "variable", "type", "bin_number", "bin_range",
                "event_count", "nonevent_count", "total_count", "event_rate",
                "woe", "iv_contrib", "total_iv", "importance", "model_type"
            ]
            df = df[column_order]
            
            # Format numeric columns
            numeric_cols = ["event_rate", "woe", "iv_contrib", "total_iv", "importance"]
            for col in numeric_cols:
                if col in df.columns:
                    df[col] = df[col].apply(lambda x: round(x, 4) if pd.notna(x) else x)
        
        return df
    
    def build_shap_report(self, shap_values, shap_summary: Dict, final_vars: List[str]) -> pd.DataFrame:
        """Build SHAP importance report"""
        rows = []
        
        # If we have SHAP summary (mean absolute SHAP values)
        if shap_summary:
            for var in final_vars:
                if var in shap_summary:
                    rows.append({
                        "rank": 0,  # Will be updated after sorting
                        "variable": var,
                        "shap_importance": shap_summary[var],
                        "relative_importance": 0  # Will be calculated
                    })
        
        # If we have raw SHAP values, calculate additional statistics
        if shap_values is not None and hasattr(shap_values, 'values'):
            shap_array = shap_values.values
            if len(shap_array.shape) == 2:  # Binary classification
                for i, var in enumerate(final_vars):
                    if i < shap_array.shape[1]:
                        values = shap_array[:, i]
                        if not rows:  # If summary wasn't provided
                            rows.append({
                                "rank": 0,
                                "variable": var,
                                "shap_importance": np.abs(values).mean(),
                                "relative_importance": 0
                            })
                        # Add additional statistics
                        for row in rows:
                            if row["variable"] == var:
                                row["shap_std"] = np.std(values)
                                row["shap_min"] = np.min(values)
                                row["shap_max"] = np.max(values)
                                row["shap_median"] = np.median(values)
                                break
        
        df = pd.DataFrame(rows)
        
        if not df.empty:
            # Sort by SHAP importance (descending)
            df = df.sort_values("shap_importance", ascending=False).reset_index(drop=True)
            
            # Calculate relative importance
            max_importance = df["shap_importance"].max()
            if max_importance > 0:
                df["relative_importance"] = (df["shap_importance"] / max_importance * 100).round(2)
            
            # Update rank
            df["rank"] = range(1, len(df) + 1)
            
            # Reorder columns
            column_order = ["rank", "variable", "shap_importance", "relative_importance"]
            if "shap_std" in df.columns:
                column_order.extend(["shap_std", "shap_min", "shap_max", "shap_median"])
            
            df = df[column_order]
            
            # Format numeric columns
            numeric_cols = ["shap_importance", "shap_std", "shap_min", "shap_max", "shap_median"]
            for col in numeric_cols:
                if col in df.columns:
                    df[col] = df[col].apply(lambda x: round(x, 6) if pd.notna(x) else x)
        
        return df
    
    def _build_metadata_report(self) -> pd.DataFrame:
        """Build metadata report"""
        meta_data = {
            "run_id": self.cfg.run_id,
            "timestamp": datetime.now().isoformat(),
            "random_state": self.cfg.random_state,
            "cv_folds": self.cfg.cv_folds,
            "hpo_timeout_sec": getattr(self.cfg, 'hpo_timeout_sec', getattr(self.cfg, 'optuna_timeout', None)),
            "hpo_trials": getattr(self.cfg, 'hpo_trials', getattr(self.cfg, 'n_trials', 100)),
            "enable_dual_pipeline": self.cfg.enable_dual_pipeline,
            "psi_threshold": self.cfg.psi_threshold,
            "iv_min": self.cfg.iv_min,
            "rho_threshold": self.cfg.rho_threshold
        }
        
        rows = [{"key": k, "value": str(v)} for k, v in meta_data.items()]
        return pd.DataFrame(rows)